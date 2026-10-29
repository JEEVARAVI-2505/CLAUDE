"""XLSX writer with typed cells; text is stored as string cells so it can never become a formula."""

from __future__ import annotations

import os
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .export_common import clean_text
from .export_csv import translate_os_error

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")


def _put(ws, row: int, col: int, value: Any, kind: str) -> None:
    cell = ws.cell(row=row, column=col)
    if value is None:
        return
    if isinstance(value, str):
        cell.value = clean_text(value)
        cell.data_type = "s"  # force a string even if it starts with '=' (no formula is ever written)
        return
    cell.value = value
    if isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"
    elif isinstance(value, Decimal):
        cell.number_format = "#,##0.00" if kind == "money" else "0.##" if kind in ("qty", "rate") else "General"


def add_sheet(wb: Workbook, title: str, headers: list[str], keys: list[str], kinds: list[str], rows: list[dict[str, Any]], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c)
        cell.value, cell.data_type = clean_text(h), "s"
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, 2):
        for c, (k, kind) in enumerate(zip(keys, kinds), 1):
            _put(ws, r, c, row.get(k), kind)
    for c, h in enumerate(headers, 1):
        width = max([len(h)] + [len(str(row.get(keys[c - 1]) or "")) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 48)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def write_xlsx(path: Path, sheets: list[tuple[str, list[str], list[str], list[str], list[dict[str, Any]]]]) -> None:
    wb = Workbook()
    for i, (title, headers, keys, kinds, rows) in enumerate(sheets):
        add_sheet(wb, title, headers, keys, kinds, rows, first=(i == 0))
    path = Path(path)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=path.name + ".", suffix=".part")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except OSError as e:
        Path(tmp).unlink(missing_ok=True)
        raise translate_os_error(e, path) from e
