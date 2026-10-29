import csv
import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from documind.domain.exceptions import ApprovalBlockedError, ExportError, InvalidValueError
from documind.domain.models import DocKind, DocStatus
from documind.samples import generate, write_sample
from tests.helpers import make_ctx, process_all, sample_pdf


@pytest.fixture()
def ctx(tmp_home):
    c = make_ctx(); yield c; c.shutdown()


def processed(ctx, tmp, seed=1, layout="gst_classic"):
    d, p = sample_pdf(tmp, seed, layout, name=f"r{seed}.pdf")
    return d, process_all(ctx, [ctx.docs.import_one(p).doc_id])[0]


def test_correction_preserves_source_and_audits(ctx, tmp_home):
    d, rec = processed(ctx, tmp_home)
    orig = rec.get("supplier_name")
    r = ctx.review.correct_field(rec.id, "supplier_name", "  Corrected   Traders  ")
    f = r.fields["supplier_name"]
    assert f.value == "Corrected Traders" and f.extracted_value == orig and f.raw and f.origin == "corrected" and f.state == "corrected"
    trail = ctx.repo.audit_trail(rec.id)
    assert trail[-1]["action"] == "corrected" and trail[-1]["old_value"] == orig and trail[-1]["new_value"] == "Corrected Traders"
    back = ctx.review.revert_field(rec.id, "supplier_name").fields["supplier_name"]
    assert back.value == orig and back.state == "extracted"


def test_invalid_input_rejected_with_clear_message(ctx, tmp_home):
    _, rec = processed(ctx, tmp_home)
    with pytest.raises(InvalidValueError) as e:
        ctx.review.correct_field(rec.id, "invoice_date", "31/02/2025")
    assert "valid date" in e.value.user_message
    with pytest.raises(InvalidValueError):
        ctx.review.correct_field(rec.id, "total", "twelve")
    with pytest.raises(InvalidValueError):
        ctx.review.correct_field(rec.id, "currency", "rupees")
    with pytest.raises(InvalidValueError):
        ctx.review.correct_field(rec.id, "__evil__", "x")


def test_correcting_total_triggers_revalidation(ctx, tmp_home):
    _, rec = processed(ctx, tmp_home)
    assert rec.error_count == 0
    r = ctx.review.correct_field(rec.id, "total", "5000")
    assert "TOTAL_MISMATCH" in {i.code for i in r.issues}
    r2 = ctx.review.correct_cell(rec.id, 0, "line_total", "1.00")
    assert "LINE_MISMATCH" in {i.code for i in r2.issues}


def test_approval_flow_blocks_errors_until_override_and_locks_editing(ctx, tmp_home):
    _, rec = processed(ctx, tmp_home)
    ctx.review.correct_field(rec.id, "total", "5000")
    with pytest.raises(ApprovalBlockedError):
        ctx.review.approve(rec.id)
    with pytest.raises(ApprovalBlockedError):
        ctx.review.approve(rec.id, override_reason="   ")
    ap = ctx.review.approve(rec.id, override_reason="Supplier confirmed the printed total is a typo")
    assert ap.status == DocStatus.APPROVED and all(f.confirmed for f in ap.fields.values() if f.value) and ap.fields["total"].state == "validated"
    assert any("override" in (a["note"] or "") for a in ctx.repo.audit_trail(rec.id))
    with pytest.raises(ApprovalBlockedError):
        ctx.review.correct_field(rec.id, "supplier_name", "x")
    re = ctx.review.reopen(rec.id)
    assert re.status == DocStatus.NEEDS_REVIEW and not any(f.confirmed for f in re.fields.values())


def test_line_add_delete_and_kind_change(ctx, tmp_home):
    _, rec = processed(ctx, tmp_home)
    n = len(rec.items)
    assert len(ctx.review.add_line(rec.id).items) == n + 1
    assert len(ctx.review.delete_line(rec.id, n).items) == n
    assert ctx.review.set_kind(rec.id, DocKind.CREDIT_NOTE).kind == DocKind.CREDIT_NOTE


# ---- export --------------------------------------------------------------------------------
def approve_all(ctx, recs):
    for r in recs:
        ctx.review.approve(r.id)


def test_xlsx_export_typed_values_and_roundtrip(ctx, tmp_home):
    d, rec = processed(ctx, tmp_home, 5)
    approve_all(ctx, [rec])
    out = tmp_home / "o.xlsx"
    res = ctx.exports.export(out, fmt="xlsx")
    assert res.documents == 1 and res.line_rows == len(d.items)
    wb = load_workbook(out)
    ws = wb["Documents"]; hdr = [c.value for c in ws[1]]; row = {h: c.value for h, c in zip(hdr, ws[2])}
    assert isinstance(row["Document date"], __import__("datetime").datetime) and row["Document date"].date() == d.doc_date
    assert Decimal(str(row["Total"])) == Decimal(d.truth["total"]) and row["Review status"] == "approved" and row["Source file"] == "r5.pdf"
    assert ws.cell(2, hdr.index("Total") + 1).number_format == "#,##0.00"
    li = wb["Line items"]; lh = [c.value for c in li[1]]
    assert [Decimal(str(li.cell(i + 2, lh.index("Line total") + 1).value)) for i in range(len(d.items))] == [i.line_total for i in d.items]
    assert {"Validation", "About"} <= set(wb.sheetnames)


def test_export_only_approved_by_default(ctx, tmp_home):
    _, a = processed(ctx, tmp_home, 6); _, b = processed(ctx, tmp_home, 7)
    approve_all(ctx, [a])
    res = ctx.exports.export(tmp_home / "x.csv", fmt="csv")
    assert res.documents == 1
    res2 = ctx.exports.export(tmp_home / "y.csv", fmt="csv", approved_only=False)
    assert res2.documents == 2


def test_csv_roundtrip_unicode_and_types(ctx, tmp_home):
    _, rec = processed(ctx, tmp_home, 8)
    ctx.review.correct_field(rec.id, "supplier_name", "Müller & Söhne GmbH – बिल 北京 😀")
    ctx.review.approve(rec.id, override_reason="n/a") if rec.error_count else ctx.review.approve(rec.id)
    p = tmp_home / "u.csv"; ctx.exports.export(p, fmt="csv", include_lines=False)
    raw = p.read_bytes(); assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0]["Supplier"] == "Müller & Söhne GmbH – बिल 北京 😀"
    date.fromisoformat(rows[0]["Document date"]); Decimal(rows[0]["Total"])


EVIL = ["=HYPERLINK(\"http://evil.example\",\"x\")", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1+1)", "\t=1+1", "=1+1\n=2+2", "  =spaced"]


@pytest.mark.parametrize("payload", EVIL)
def test_formula_injection_neutralised_in_csv_and_xlsx(ctx, tmp_home, payload):
    _, rec = processed(ctx, tmp_home, 9)
    ctx.review.correct_field(rec.id, "supplier_name", payload)
    ctx.review.correct_field(rec.id, "customer_name", payload)
    ctx.review.approve(rec.id, override_reason="test") if ctx.repo.get(rec.id).error_count else ctx.review.approve(rec.id)
    c = tmp_home / "i.csv"; ctx.exports.export(c, fmt="csv", include_lines=False)
    got = list(csv.DictReader(io.StringIO(c.read_text("utf-8-sig"))))[0]["Supplier"]
    assert got.startswith("'") and got.lstrip("'").strip() != "" and not got.lstrip(" ").startswith(("=", "+", "-", "@", "\t"))
    x = tmp_home / "i.xlsx"; ctx.exports.export(x, fmt="xlsx")
    ws = load_workbook(x)["Documents"]
    cell = next(c for c in ws[2] if isinstance(c.value, str) and c.value.strip().lstrip("'")[:1] in "=+-@" or (isinstance(c.value, str) and "=" in c.value[:3]))
    assert cell.data_type == "s"
    with zipfile.ZipFile(x) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        assert "<f>" not in sheet and "<f " not in sheet


def test_export_failures_are_friendly(ctx, tmp_home, monkeypatch):
    _, rec = processed(ctx, tmp_home, 10); approve_all(ctx, [rec])
    with pytest.raises(ExportError):
        ctx.exports.export(tmp_home / "nope" / "x.xlsx")
    with pytest.raises(ExportError):
        ctx.exports.export(tmp_home)                                            # a folder, not a file
    with pytest.raises(ExportError):
        ctx.exports.export(tmp_home / "z.xlsx", columns=["not_a_column"])
    import os, errno
    real = os.replace
    monkeypatch.setattr(os, "replace", lambda *a, **k: (_ for _ in ()).throw(OSError(errno.ENOSPC, "No space left on device")))
    from documind.domain.exceptions import DiskFullError
    with pytest.raises(DiskFullError):
        ctx.exports.export(tmp_home / "full.csv", fmt="csv")
    monkeypatch.setattr(os, "replace", real)
    assert not list(tmp_home.glob("full.csv*"))                                 # no half-written leftovers
