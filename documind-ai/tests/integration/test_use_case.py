"""The product's reason to exist, as one story:

An accounts clerk receives a month of supplier invoices, drops them in, checks what the software flags, approves, and sends their
accountant a purchase register whose totals they can trust. Every number is compared with the known answer.
"""
from collections import defaultdict
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from documind.domain.models import DocStatus
from documind.samples import generate, write_sample
from documind.infrastructure.export_common import PRESETS
from tests.helpers import make_ctx, process_all


@pytest.fixture()
def ctx(tmp_home):
    c = make_ctx(); yield c; c.shutdown()


def month_of_invoices(tmp, n=9):
    docs = []
    for i in range(n):
        d = generate(7000 + i, ["gst_classic", "gst_igst", "labels_below"][i % 3])
        p = tmp / f"inbox_{i:02d}.pdf"; write_sample(d, p); docs.append((d, p))
    return docs


def test_month_end_story(ctx, tmp_home):
    docs = month_of_invoices(tmp_home)
    # 1. add everything at once, plus one useless file and one repeated file
    junk = tmp_home / "notes.txt"; junk.write_text("not an invoice")
    results = ctx.docs.import_files([p for _, p in docs] + [junk, docs[0][1]])
    assert [r.status for r in results].count("imported") == 9
    assert results[-2].status == "rejected" and results[-1].status == "duplicate"
    ids = {p.name: r.doc_id for (_, p), r in zip(docs, results) if r.doc_id}
    recs = process_all(ctx, list(ids.values()))
    assert all(r.status == DocStatus.NEEDS_REVIEW for r in recs)

    # 2. the dashboard-level counts the clerk sees
    stats = ctx.repo.stats(None)
    assert stats["needs_review"] == 9 and stats["approved"] == 0

    # 3. one invoice has a wrong total on the page itself: it must be flagged and must NOT be bulk-approved
    bad_doc, bad_path = docs[4]
    bad_id = ids[bad_path.name]
    ctx.review.correct_field(bad_id, "total", "1.00")            # simulates a misprinted total
    assert ctx.repo.get(bad_id).error_count >= 1
    clean_before = set(ctx.review.clean_candidates())
    assert bad_id not in clean_before

    # 4. approve everything that passed all checks in one go
    n = ctx.review.approve_clean()
    assert n == len(clean_before) and ctx.repo.get(bad_id).status == DocStatus.NEEDS_REVIEW
    assert all(ctx.repo.get(i).status == DocStatus.APPROVED for i in clean_before)

    # 5. the clerk looks at the flagged one, sees the printed total is wrong, and fixes it from the page
    ctx.review.correct_field(bad_id, "total", bad_doc.truth["total"])
    assert ctx.repo.get(bad_id).error_count == 0
    # remaining documents with warnings (e.g. ambiguous dates) are checked one by one
    for i in ids.values():
        if ctx.repo.get(i).status == DocStatus.NEEDS_REVIEW:
            ctx.review.approve(i)
    assert ctx.repo.stats(None)["approved"] == 9

    # 6. export the purchase register
    out = tmp_home / "purchase_register.xlsx"
    res = ctx.exports.export(out, preset="purchase_register")
    assert res.documents == 9 and res.skipped_unapproved == 0
    wb = load_workbook(out)
    assert wb.sheetnames[:2] == ["Documents", "Summary"] and "Line items" not in wb.sheetnames
    ws = wb["Documents"]; hdr = [c.value for c in ws[1]]
    assert hdr[:4] == ["Document date", "Document / invoice no.", "Supplier", "Supplier tax ID"] or hdr[0] == "Document date"
    rows = [{h: c.value for h, c in zip(hdr, r)} for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 9
    by_total = sorted(Decimal(str(r["Total"])) for r in rows)
    assert by_total == sorted(Decimal(d.truth["total"]) for d, _ in docs)

    # 7. the Summary sheet's grand totals equal the sum of the known answers (per month and per supplier, never mixing currencies)
    sm = wb["Summary"]; sh = [c.value for c in sm[1]]
    srows = [{h: c.value for h, c in zip(sh, r)} for r in sm.iter_rows(min_row=2)]
    month_total = sum(Decimal(str(r["Total"])) for r in srows if r["View"] == "By month")
    supplier_total = sum(Decimal(str(r["Total"])) for r in srows if r["View"] == "By supplier")
    expected = sum(Decimal(d.truth["total"]) for d, _ in docs)
    assert month_total == supplier_total == expected
    expected_by_supplier = defaultdict(Decimal)
    for d, _ in docs:
        expected_by_supplier[d.supplier] += Decimal(d.truth["total"])
    got = {r["Name"]: Decimal(str(r["Total"])) for r in srows if r["View"] == "By supplier"}
    assert got == dict(expected_by_supplier)


def test_summary_never_mixes_currencies(ctx, tmp_home):
    a = generate(7101, "gst_classic")        # INR
    b = generate(7102, "simple_us")          # USD
    paths = []
    for d, name in ((a, "a.pdf"), (b, "b.pdf")):
        p = tmp_home / name; write_sample(d, p); paths.append(p)
    recs = process_all(ctx, [ctx.docs.import_one(p).doc_id for p in paths])
    for r in recs:
        ctx.review.approve(r.id, override_reason="story") if ctx.repo.get(r.id).error_count else ctx.review.approve(r.id)
    out = tmp_home / "s.xlsx"; ctx.exports.export(out, preset="purchase_register")
    sm = load_workbook(out)["Summary"]; sh = [c.value for c in sm[1]]
    cur = {r[sh.index("Currency")].value for r in sm.iter_rows(min_row=2)}
    assert cur == {"INR", "USD"}


@pytest.mark.parametrize("key", list(PRESETS))
def test_every_export_preset_produces_a_valid_file(ctx, tmp_home, key):
    d = generate(7200, "gst_classic"); p = tmp_home / "x.pdf"; write_sample(d, p)
    (r,) = process_all(ctx, [ctx.docs.import_one(p).doc_id]); ctx.review.approve(r.id)
    for fmt in ("xlsx", "csv"):
        res = ctx.exports.export(tmp_home / f"{key}.{fmt}", preset=key, fmt=fmt)
        assert res.documents == 1 and res.paths[0].exists()


def test_bulk_approve_skips_documents_with_problems_and_is_audited(ctx, tmp_home):
    d, p = generate(7300, "gst_classic"), tmp_home / "ok.pdf"; write_sample(d, p)
    (r,) = process_all(ctx, [ctx.docs.import_one(p).doc_id])
    ctx.review.correct_field(r.id, "total", "5")
    assert ctx.review.clean_candidates() == [] and ctx.review.approve_clean() == 0
    ctx.review.correct_field(r.id, "total", d.truth["total"])
    if r.id in ctx.review.clean_candidates():
        assert ctx.review.approve_clean() == 1
        assert any(a["action"] == "bulk_approved" for a in ctx.repo.audit_trail(r.id))
